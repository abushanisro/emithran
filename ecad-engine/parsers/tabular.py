"""Tabular customer files (Pick Place CSV or workbook, BOM workbook) as header + rows.

The header row is found, not assumed: Altium Pick Place CSVs carry a dozen lines of design information
before it, and they are GBK encoded rather than UTF-8.
"""
from __future__ import annotations

import csv
import io
import struct
import zipfile
from dataclasses import dataclass, field

from openpyxl import load_workbook

ENCODINGS = ("utf-8-sig", "gbk")

# What a customer workbook may expand to (openpyxl inflates its parts in memory). A BOM or Pick Place sheet is well under this.
XLSX_MAX_ENTRIES = 5000
XLSX_MAX_ENTRY_BYTES = 50 * 1024 * 1024
XLSX_MAX_TOTAL_BYTES = 50 * 1024 * 1024
XLSX_MAX_RATIO = 100


@dataclass
class Table:
    headers: list[str]
    rows: list[dict[str, str]]
    preamble: list[str] = field(default_factory=list)


def _decode(data: bytes) -> str:
    for encoding in ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("file is neither UTF-8 nor GBK text")


def read_csv_table(data: bytes, header_marker: str) -> Table:
    """The table starts at the first line whose first cell equals `header_marker`."""
    lines = _decode(data).splitlines()
    start = next((i for i, line in enumerate(lines) if (next(csv.reader([line]), None) or [""])[0] == header_marker), None)
    if start is None:
        raise ValueError(f"no header row starting with {header_marker!r}")
    reader = csv.DictReader(io.StringIO("\n".join(lines[start:])))
    rows = [{k: (v or "").strip() for k, v in row.items() if k is not None} for row in reader]
    return Table(headers=list(reader.fieldnames or []), rows=rows, preamble=lines[:start])


def check_workbook_archive(data: bytes) -> None:
    """Measure an .xlsx from its zip central directory before openpyxl inflates any part of it.

    The same rules as the backend's inspect-zip.ts: a bounded entry count, each part and the total declared size within
    XLSX_MAX_*_BYTES, a total compression ratio of at most XLSX_MAX_RATIO once past 1 MB, and every entry owning its own bytes
    (local header and data inside the file, not shared with or overlapping another entry's), so the declared sizes bound what
    inflating can produce. Raises ValueError, which the service reports as an unreadable file.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
    except (zipfile.BadZipFile, zipfile.LargeZipFile, OSError) as error:
        raise ValueError("workbook is not a readable zip") from error
    if len(infos) > XLSX_MAX_ENTRIES:
        raise ValueError(f"workbook holds more than {XLSX_MAX_ENTRIES} parts")
    total = 0
    ranges: list[tuple[int, int]] = []
    for info in infos:
        if info.file_size > XLSX_MAX_ENTRY_BYTES:
            raise ValueError("a workbook part expands beyond the allowed size")
        total += info.file_size
        if total > XLSX_MAX_TOTAL_BYTES:
            raise ValueError("workbook expands beyond the allowed size")
        start = info.header_offset
        header = data[start : start + 30]
        if len(header) < 30 or header[:4] != b"PK\x03\x04":
            raise ValueError("a workbook part's local header is missing")
        name_len, extra_len = struct.unpack("<HH", header[26:30])
        end = start + 30 + name_len + extra_len + info.compress_size
        if end > len(data):
            raise ValueError("a workbook part's data is truncated")
        ranges.append((start, end))
    ranges.sort()
    if any(start < previous_end for (_, previous_end), (start, _) in zip(ranges, ranges[1:])):
        raise ValueError("workbook parts share or overlap their data")
    compressed = sum(info.compress_size for info in infos)
    if compressed > len(data):
        raise ValueError("workbook parts share or overlap their data")
    if total > 1024 * 1024 and total / max(compressed, 1) > XLSX_MAX_RATIO:
        raise ValueError("workbook compression ratio is implausibly high")


def read_sheet_table(data: bytes, header_marker: str | None = None, sheet: str | None = None) -> Table:
    """header_marker None: the header is the first non-empty row."""
    check_workbook_archive(data)
    book = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    worksheet = book[sheet] if sheet else book.worksheets[0]
    cells = [["" if c is None else str(c).strip() for c in row] for row in worksheet.iter_rows(values_only=True)]
    start = next((i for i, row in enumerate(cells) if any(row) and (header_marker is None or row[0] == header_marker)), None)
    if start is None:
        raise ValueError(f"no header row starting with {header_marker!r}")
    headers = cells[start]
    rows = [dict(zip(headers, row)) for row in cells[start + 1 :] if any(row)]
    return Table(headers=headers, rows=rows, preamble=[" ".join(r) for r in cells[:start]])
